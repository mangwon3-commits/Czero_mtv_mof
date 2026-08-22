# -*- coding: utf-8 -*-
"""학술제 1매 요약 포스터 — 편집용 .pptx (A4 세로 슬라이드 1장).

PDF 판과 같은 내용·팔레트. 모든 요소가 개별 도형이라 위치·글자를 자유롭게
고칠 수 있다. VESTA/ParaView 자리는 흰 사각형으로 비워 두었다.
규정: 1매 · 모든 텍스트 10pt 이상 · 블라인드(팀·개인 식별 정보 금지).
"""
import os
from pptx import Presentation
from pptx.util import Cm, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

OUT = os.path.expanduser('~/mof_project/04_Analysis/학술제_요약포스터.pptx')
FIG = os.path.expanduser('~/mof_project/04_Analysis/figs/')

NAVY   = RGBColor(0x1F, 0x3B, 0x57)
ORANGE = RGBColor(0xE0, 0x8A, 0x3C)
SAGE   = RGBColor(0x8F, 0xA3, 0x7E)
SAGE_L = RGBColor(0xE8, 0xEE, 0xE1)
GREY   = RGBColor(0x6A, 0x6A, 0x6A)
GREY_L = RGBColor(0xC8, 0xC8, 0xC8)
INK    = RGBColor(0x22, 0x26, 0x2A)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
FONT   = '맑은 고딕'

prs = Presentation()
prs.slide_width, prs.slide_height = Cm(21.0), Cm(29.7)   # A4 세로
s = prs.slides.add_slide(prs.slide_layouts[6])           # 빈 레이아웃

M  = 1.0     # 여백 cm
CW = (21.0 - 2*M - 0.6) / 2
LX, RX = M, M + CW + 0.6


def box(x, y, w, h, fill=None, line=None, dash=False):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(0.75)
        if dash:
            from pptx.enum.dml import MSO_LINE_DASH_STYLE
            sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    sh.text_frame.text = ''
    return sh


def text(x, y, w, h, runs, size=10, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, space=2, line_sp=1.06):
    """runs: 문자열 또는 [(글자, 굵기), ...]"""
    tb = s.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_sp
    p.space_after = Pt(space)
    if isinstance(runs, str):
        runs = [(runs, bold)]
    for t, b in runs:
        r = p.add_run(); r.text = t
        r.font.size = Pt(size); r.font.bold = b
        r.font.color.rgb = color; r.font.name = FONT
        r.font._rPr.set('altLang', 'ko-KR')
    return tb


def para(tb, runs, size=10, color=INK, align=PP_ALIGN.LEFT, space=2, line_sp=1.06):
    """기존 텍스트박스에 문단 추가."""
    p = tb.text_frame.add_paragraph()
    p.alignment = align; p.line_spacing = line_sp; p.space_after = Pt(space)
    if isinstance(runs, str):
        runs = [(runs, False)]
    for t, b in runs:
        r = p.add_run(); r.text = t
        r.font.size = Pt(size); r.font.bold = b
        r.font.color.rgb = color; r.font.name = FONT
    return p


def header(x, y, w, label):
    box(x, y, w, 0.72, fill=NAVY)
    text(x + 0.32, y + 0.10, w - 0.4, 0.55, label, size=12.5, color=WHITE,
         bold=True)
    return y + 0.72 + 0.18


def caption(x, y, w, s_, lines=2):
    h = 0.44 * lines
    text(x, y, w, h, s_, size=10, color=GREY, line_sp=1.02)
    return y + h + 0.18


def table(x, y, w, rows, colw, size=10):
    """머리행 + 본문. 높이를 돌려준다."""
    rh = 0.52
    t = s.shapes.add_table(len(rows), len(rows[0]), Cm(x), Cm(y),
                           Cm(w), Cm(rh*len(rows))).table
    for j, cwd in enumerate(colw):
        t.columns[j].width = Cm(cwd)
    for i, row in enumerate(rows):
        t.rows[i].height = Cm(rh)
        for j, v in enumerate(row):
            cell = t.cell(i, j)
            cell.margin_left = cell.margin_right = Cm(0.12)
            cell.margin_top = cell.margin_bottom = Emu(0)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = SAGE_L if i == 0 else WHITE
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT
            r = p.add_run(); r.text = str(v)
            r.font.size = Pt(size); r.font.bold = (i == 0)
            r.font.color.rgb = INK; r.font.name = FONT
    return y + rh*len(rows) + 0.25

# ══════════════════ 제목 ══════════════════
box(0, 0, 21.0, 2.0, fill=NAVY)
box(0, 2.0, 21.0, 0.15, fill=ORANGE)
text(0.8, 0.30, 19.4, 0.9, 'CO₂ 포집을 위한 MTV-ZIF의 전산 설계; ZIF-69를 중심으로',
     size=17, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
text(0.8, 1.24, 19.4, 0.6,
     '습윤 배가스(CO₂ 0.15 bar · 298 K · RH 90%)에서 최적 치환율을 찾는 6단계 평가 파이프라인',
     size=10.5, color=RGBColor(0xD6, 0xE0, 0xEA), align=PP_ALIGN.CENTER)

Y0 = 2.75
BOT = 25.9          # 결론 띠 상단

# ══════════════════ 왼쪽 단 ══════════════════
y = header(LX, Y0, CW, '1.  연구 질문')
tb = text(LX, y, CW, 1.6,
          '실제 배가스에는 수분이 동반되므로 건조 조건의 CO₂ 흡착량만으로는 포집 소재를 '
          '평가할 수 없다. ZIF 유기 링커를 부분 치환한 다변량(MTV) 골격을 설계하고 '
          '다음 한 질문에 답한다.', size=10)
y += 1.70

box(LX, y, CW, 0.95, fill=SAGE_L)
box(LX, y, 0.16, 0.95, fill=SAGE)
text(LX + 0.42, y + 0.20, CW - 0.6, 0.6,
     '습윤 배가스에서 어느 치환율이 가장 유리한가?', size=11, color=NAVY, bold=True)
y += 1.22

text(LX, y, CW, 1.4,
     '치환기가 늘면 정전기 자리가 늘어 CO₂ 흡착이 오르지만(가설 1), 같은 자리가 물도 '
     '잡는다(가설 2). 판정 기준을 계산 전에 등록하고 측정으로 가린다.', size=10)
y += 1.42

# ── VESTA 자리 ──
box(LX, y, CW, 3.8, fill=WHITE, line=GREY_L, dash=True)
y += 3.8 + 0.16
y = caption(LX, y, CW,
            '그림 1. ZIF-69(gme) 단위셀과 치환 가능한 24개 -Cl 자리. '
            '벤조이미다졸레이트 b2 자리가 공동을 향한다.')

y = header(LX, y, CW, '2.  계산 방법')
text(LX, y, CW, 1.4,
     '흡착 RASPA2 대정준 몬테카를로 · 헨리 상수와 흡착열 Widom 삽입 · 이완 GFN-FF'
     '(셀 고정) · 전하 PACMAN DDEC6 · 기공 지표 Zeo++. 설정은 전 계산에서 고정했다.',
     size=10)
y += 1.44
y = table(LX, y, CW, [
    ['항목', '설정'],
    ['사이클', '초기화 5,000 + 생산 15,000'],
    ['골격 힘장 / 전하', 'UFF_MOF / PACMAN DDEC6'],
    ['CO₂ · 물 모델', 'García-Sánchez 2009 · TIP5P-Ew'],
    ['정전기 / 컷오프', 'Ewald 1e-6 / 12 Å'],
    ['재생', 'TSA 373 K / VSA 0.05 bar'],
], [CW*0.38, CW*0.62])

y = header(LX, y, CW, '3.  작용기 선정')
text(LX, y, CW, 1.4,
     '작용기 7종 중 -SO₃H만 목표대(30~40)에 든다. -OH를 -CH₃로 막은 -SO₂CH₃와 같은 '
     '치환율에서 2.41 kJ/mol 벌어지며, 그 차이가 수소결합 공여 -OH의 몫이다.', size=10)
y += 1.44
y = table(LX, y, CW, [
    ['작용기 (최적 조성)', 'Q_st', '무치환 대비'],
    ['-SO₃H (100%)', '34.01', '+11.59'],
    ['-SO₂CH₃ (75%)', '29.01', '+6.59'],
    ['-NO₂ (100%)', '28.87', '+6.45'],
    ['-Br (100%)', '22.73', '+0.31'],
    ['무치환 모체', '22.42', '기준'],
], [CW*0.46, CW*0.27, CW*0.27])
print(f'  왼쪽 단 끝 {y:.2f} cm  (한계 {BOT})')

# ══════════════════ 오른쪽 단 ══════════════════
from PIL import Image
y = header(RX, Y0, CW, '4.  결과')
FW = CW * 0.94
FX = RX + (CW - FW) / 2


def picture(y, name, cap, lines=2):
    iw, ih = Image.open(FIG + name).size
    h = FW * ih / iw
    s.shapes.add_picture(FIG + name, Cm(FX), Cm(y), width=Cm(FW), height=Cm(h))
    return caption(RX, y + h + 0.16, CW, cap, lines=lines)


y = picture(y, 'fig2_ladder.png',
            '그림 2. 치환율에 따른 등량 흡착열과 0.15 bar 로딩. '
            '흡착열 22.4→34.0 kJ/mol, 로딩 2.6배. 오차막대 1σ.')
y = picture(y, 'fig3_water.png',
            '그림 3. RH90 유지율(좌축)과 H₂O/CO₂ 비(우축). 유지율은 25~75%에서 '
            '고원을 이루고 100%에서 꺾인다 — 점진적 악화가 아닌 문턱 현상.')
y = picture(y, 'fig4_windows.png',
            '그림 4. 두 관문의 창. 안정성 창은 치환 자리 하나(58.3→62.5%) 사이에서 '
            '닫히고, 흡착 창은 66.7→75%에서 열린다. 회색이 두 창 사이의 간격.')

# ── ParaView 자리 ──
box(RX, y, CW, 1.9, fill=WHITE, line=GREY_L, dash=True)
y += 1.9 + 0.16
y = caption(RX, y, CW,
            '그림 5. 전하 ON/OFF 차분 밀도맵 — 정전기가 CO₂를 끌어들인 자리 '
            '(정전기 기여 무치환 41% → 고치환 69%).')
print(f'  오른쪽 단 끝 {y:.2f} cm  (한계 {BOT})')

# ══════════════════ 결론 띠 ══════════════════
box(0, BOT, 21.0, 29.7 - BOT, fill=SAGE_L)
box(0, BOT, 21.0, 0.14, fill=NAVY)
text(M, BOT + 0.30, 4.0, 0.6, '결론', size=12.5, color=NAVY, bold=True)

BW = 7.2
box(21.0 - M - BW, BOT + 0.85, BW, 2.35, fill=NAVY)
text(21.0 - M - BW, BOT + 1.02, BW, 0.6, '-SO₃H  58.3 % (14/24 자리)',
     size=12.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
text(21.0 - M - BW, BOT + 1.72, BW, 0.5, '습윤 TSA 작업 용량 0.833 ± 0.025 mol/kg',
     size=10, color=RGBColor(0xDC, 0xE5, 0xEE), align=PP_ALIGN.CENTER)
text(21.0 - M - BW, BOT + 2.30, BW, 0.5, '무치환 대비 +77 % · 차순위와 2.10σ',
     size=10, color=RGBColor(0xDC, 0xE5, 0xEE), align=PP_ALIGN.CENTER)

tw = 21.0 - 2*M - BW - 0.5
tb = text(M, BOT + 0.92, tw, 2.4,
          '① 건조 성능 1위(100% 치환)가 습윤에서 꼴찌다. CO₂를 잡는 -OH 자리가 물도 '
          '잡기 때문이며, 최대 치환은 최적이 아니다.', size=10, line_sp=1.04)
para(tb, '② 안정성 창(≤58.3%)과 흡착 목표대 창(≥75%)이 겹치지 않아, -SO₃H 단독으로는 '
         '두 관문을 동시에 넘는 조성이 없다.', size=10, line_sp=1.04)
para(tb, '③ 재생의 지배항은 흡착열이 아니라 현열(2~5배)이며, 이득은 작업 용량으로 '
         '난방비를 희석한 데서 온다.', size=10, line_sp=1.04)

prs.save(OUT)
print('저장:', OUT)
